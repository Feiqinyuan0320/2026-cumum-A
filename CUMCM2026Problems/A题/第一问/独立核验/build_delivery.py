"""Build a consistent writer handoff from completed independent calculations."""
from pathlib import Path
import json,hashlib
import numpy as np
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from verify_q1 import ROOT

DEST=ROOT.parent/'核验后交付'
TIMES=[100,300,600,900,1200,1500,1800]

def table(values):
    lines=['| 时间/s | 0 cm | 0.5 cm | 1.0 cm | 1.5 cm | 2.0 cm |',
           '| ---: | ---: | ---: | ---: | ---: | ---: |']
    for t in TIMES:
        lines.append('| '+str(t)+' | '+' | '.join(f'{x:.4f}' for x in values[t,::5])+' |')
    return '\n'.join(lines)

def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    DEST.mkdir(exist_ok=True)
    z=np.load(ROOT/'runs/bdf6400_128.npz')
    T,C=z['T'],z['C']
    audit=json.loads((ROOT/'核验数据.json').read_text())
    meta=json.loads((ROOT/'runs/bdf6400_128.json').read_text())
    wb=openpyxl.Workbook();wb.remove(wb.active)
    for name,key in [('温度','T'),('水分浓度','C')]:
        ws=wb.create_sheet(name)
        ws.append(['时间/s']+[round(j/10,1) for j in range(21)])
        for t in range(1,1801):ws.append([t]+[float(f'{v:.4f}') for v in z[key][t]])
        ws.freeze_panes='B2'
        for cell in ws[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='284B63')
            cell.alignment=Alignment(horizontal='center')
        for row in ws.iter_rows(min_row=2,min_col=2):
            for cell in row:cell.number_format='0.0000'
        for cell in list(ws[1])[1:]:cell.number_format='0.0'
        ws.column_dimensions['A'].width=12
        for j in range(2,23):ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width=11
    excel=DEST/'result1.xlsx';wb.save(excel)
    checked=openpyxl.load_workbook(excel,data_only=True)
    for name,key in [('温度','T'),('水分浓度','C')]:
        s=checked[name]
        assert (s.max_row,s.max_column)==(1801,22)
        vals=np.array(list(s.values)[1:],float)
        assert np.array_equal(vals[:,0],np.arange(1,1801))
        assert np.array_equal(vals[:,1:],np.round(z[key][1:],4))
        assert all(c.number_format=='0.0000' for row in s.iter_rows(min_row=2,min_col=2) for c in row)
    checked.close()
    analysis=r'''## 问题一的分析

问题一给定药材的几何尺寸、初始状态、物性参数及烘房环境随时间的变化，要求计算前1800 s内不同径向位置的温度和干基含水率。计算对象是药材内部的瞬态分布，附件1中的31个环境采样点用于确定外部边界随时间的变化。建模需要连接外界空气、药材表面和内部各位置之间的传递过程。

对温度场，外界空气先通过表面对流向药材传热，再通过内部导热影响中心区域。以半径为特征长度，热Biot数为 $h_TR/k\approx1.3889$，表面对流热阻和内部导热热阻均需在模型中保留。对水分场，表面传质与内部扩散共同影响水分迁移，且扩散系数 $D(C)$ 随局部含水率变化，求解时需要更新各位置的扩散能力。

本问采用固定半径和常数热物性，忽略潜热反馈后，两个场可分别计算。数值处理的关键是保留圆柱几何权重、正确处理轴线对称条件和表面对流边界，并对非线性水分方程进行迭代。题目同时要求论文中的35个取值点和逐秒、每0.1 cm的完整结果，因此误差检查应覆盖全部输出位置，不能只依据论文表格判断早期近表面的分辨率是否充分。

'''
    original=(ROOT.parent/'第一问正文.md').read_text()
    model=original.split('## 3. 有限体积离散框架')[0]
    model=model.replace('## 1. 模型建立',analysis+'## 1. 模型建立',1)
    numerical=r'''## 3. 数值求解

采用均匀单元中心有限体积网格。设 $\Delta r=R/N$，代表节点为 $r_i=(i-1/2)\Delta r$，控制体边界为 $r_{i-1/2}=(i-1)\Delta r$ 和 $r_{i+1/2}=i\Delta r$。控制体体积及界面面积分别为

$$
V_i=\pi L(r_{i+1/2}^2-r_{i-1/2}^2),\qquad
A_{i\pm1/2}=2\pi Lr_{i\pm1/2}.
$$

离散未知量按单元中心点值解释，以其近似控制体内的场量。定义径向向外的热通量 $q=-kT_r$ 和约化水分通量 $f=-D(C)C_r$。对任一内部界面，使用相邻节点差分近似梯度，水分界面扩散系数取调和平均

$$
D_{i+1/2}=\frac{2D(C_i)D(C_{i+1})}{D(C_i)+D(C_{i+1})}.
$$

相邻控制体共用同一界面通量，保证内部交换量在整体求和时抵消。时间离散采用二阶后向差分格式（BDF2），第一步用后向欧拉启动。用 $\phi$ 统一表示温度或含水率，$m_T=\rho c_p$、$m_C=1$，$F_T=q$、$F_C=f$，则从第二步开始有

$$
m_\phi V_i\frac{3\phi_i^{n+1}-4\phi_i^n+\phi_i^{n-1}}{2\Delta t}
=A_{i-1/2}F_{i-1/2}^{n+1}-A_{i+1/2}F_{i+1/2}^{n+1}.
$$

### 3.1 轴线与表面的离散

轴心控制体内界面面积为零，内侧通量项自然消失。输出轴线值时，在前两个中心点之间按对称偶二次函数重构：

$$
\phi(0,t)\approx\frac{9\phi_1(t)-\phi_2(t)}8.
$$

该式对应中心点值解释，不能与严格的圆环体积平均值重构混用。内部输出点采用相邻中心点之间的线性插值，表面输出值由Robin条件计算。

最外层单元中心 $r_P=R-\Delta r/2$ 与表面不在同一位置。考虑半单元的圆柱传递阻力，定义

$$
\ell_R=R\ln\frac{R}{r_P}.
$$

热边界采用

$$
q_R=\frac{T_P-T_\infty}{\ell_R/k+1/h_T},
\qquad T_s=T_\infty+\frac{q_R}{h_T}.
$$

水分边界保留半单元内扩散系数的变化，用

$$
\overline D_R=
\frac{D(C_P)+4D((C_P+C_s)/2)+D(C_s)}6
$$

近似含水率区间上的扩散系数积分均值，构造

$$
f_R=\frac{C_P-C_\infty}{\ell_R/\overline D_R+1/h_m},
\qquad C_s=C_\infty+\frac{f_R}{h_m}.
$$

这一边界重构来自局部准稳态通量关系

$$
\int_{C_s}^{C_P}D(c)\,\mathrm dc
\approx \ell_R h_m(C_s-C_\infty),
$$

其中积分均值采用Simpson近似。该关系作为半单元的空间离散处理，仍以瞬态控制体守恒式推进内部场量。

### 3.2 非线性求解与计算设置

温度方程离散后为线性三对角方程组。水分方程在每个时间步采用Picard迭代：以上一时间层含水率作为初值，根据当前迭代值更新内部界面及表面的扩散系数，解三对角线性方程组，并更新表面含水率。当全部单元和表面含水率的相邻迭代最大差小于 $10^{-10}$ 时停止迭代，再用更新后的系数检查离散方程与边界残差。温度和水分的线性子问题均采用追赶法求解。

本次结果采用 $N=6400$、$\Delta r=3.125\times10^{-6}\ \mathrm m$ 和 $\Delta t=0.0078125\ \mathrm s$。外部环境使用附件1前31个节点的分段线性插值。完整输出共1800个时刻和21个径向位置；每个场有37800个结果。

## 4. 数值核验

在全部37800个输出点上比较网格和时间步细化结果，最大绝对差如下。表中空间比较固定 $\Delta t=0.0078125\ \mathrm s$；时间比较固定 $N=3200$。

@CONVERGENCE_TABLE@

这两组比较在论文指定的35个取值点上，温度及含水率的四位小数结果均保持一致。空间细化在完整输出中仍有少数跨越舍入阈值的数值，故上述检验作为离散误差减小的证据，不据此宣称所有输出点的第四位小数均已严格确定。

温度场另采用圆柱导热的特征函数级数解进行独立验证。特征根满足 $\mu J_1(\mu)=\mathrm{Bi}_T J_0(\mu)$，在每段线性环境温度输入下解析推进模态系数。400项与800项级数结果的全输出最大差为 @SERIES_TAIL@ ℃；最终有限体积温度与800项级数结果的最大差为 @HEAT_REFERENCE_ERROR@ ℃，35个论文取值点的四位小数全部一致。

守恒检验按相应时间离散格式计算。以BDF2为例，对所有控制体求和后，温度场满足

$$
\rho c_p\sum_iV_i\frac{3T_i^{n+1}-4T_i^n+T_i^{n-1}}{2\Delta t}
=-A_Rq_R^{n+1},\qquad A_R=2\pi RL,
$$

水分场具有相同的约化守恒形式。全过程最大热量绝对平衡残差为 @HEAT_RESIDUAL@ W，约化水分绝对平衡残差为 @WATER_RESIDUAL@ $\mathrm{m^3/s}$（含水率按kg水/kg干物质计）。后者乘以恒定干物质体积密度后对应实际水质量平衡残差。最终计算每步最多进行 @PICARD@ 次Picard迭代。计算过程中未发现温度随时间下降、含水率随时间上升或径向梯度方向违反预期的情况。

上述检验用于核查离散和求解的一致性。由于缺少内部测温与含水率实测数据，不能将其解释为实验验证。

## 5. 计算结果与分析

表1给出指定时刻和位置的温度，表2给出对应的干基含水率。表格和完整结果文件使用同一计算输出，均保留四位小数。

表1　30分钟内药材的温度（℃）

@TABLE_T@

表2　30分钟内药材的水分浓度（kg/kg，干基）

@TABLE_C@

@RESULT_ANALYSIS@

## 6. 模型适用范围

本模型描述既定假设下圆柱中段的径向响应。外部水分浓度采用题目规定尺度下的等效解释；端面影响、干燥收缩及蒸发潜热对温度场的反馈均不在本问计算范围内。数值精度检验不能消除这些物理简化带来的误差。若进一步引入端面换热或表面潜热项，需要相应扩展模型并重新求解。

## 参考文献

[1] 王乐意, 李长河, 刘明政, 等. 中药材干燥技术与装备研究现状[J]. 农业工程学报, 2024, 40(2): 1–28. DOI: [10.11975/j.issn.1002-6819.202306104](https://doi.org/10.11975/j.issn.1002-6819.202306104).

[2] DA SILVA W P, E SILVA C M D P S, GAMA F J A. Estimation of thermo-physical properties of products with cylindrical shape during drying: The coupling between mass and heat[J]. Journal of Food Engineering, 2014, 141: 65–73. DOI: [10.1016/j.jfoodeng.2014.05.010](https://doi.org/10.1016/j.jfoodeng.2014.05.010).
'''
    space=audit['pairs']['bdf3200_128__bdf6400_128']
    timecheck=audit['pairs']['bdf3200_64__bdf3200_128']
    convergence='| 检验 | 对比设置 | 最大温度差/℃ | 最大含水率差/(kg/kg) |\n| --- | --- | ---: | ---: |\n'
    convergence+=f"| 空间细化 | N=3200→6400 | {space['T']['max_abs']:.6e} | {space['C']['max_abs']:.6e} |\n"
    convergence+=f"| 时间细化 | Δt=0.015625→0.0078125 s | {timecheck['T']['max_abs']:.6e} | {timecheck['C']['max_abs']:.6e} |"
    result_analysis=f'''在1800 s时，烘房温度为41.513 ℃，药材表面温度为 {T[1800,-1]:.4f} ℃，中心温度为 {T[1800,0]:.4f} ℃。表面仍低于同期烘房温度 {41.513-T[1800,-1]:.4f} ℃，表面与中心的温差为 {T[1800,-1]-T[1800,0]:.4f} ℃，说明前30分钟内药材内部尚未达到温度均匀状态。表面先响应外界升温，内部位置存在响应滞后，温度随半径增大而升高。

水分变化主要集中在表层。1800 s时，表面干基含水率由2.55降至 {C[1800,-1]:.4f} kg/kg，相对初值降低约 {(2.55-C[1800,-1])/2.55*100:.2f}%；距轴线1.5 cm处为 {C[1800,15]:.4f} kg/kg，而中心在四位小数下仍为 {C[1800,0]:.4f} kg/kg。这表示中心变化小于当前显示精度，并不表示内部水分完全没有迁移。中心与表面的含水率差为 {C[1800,0]-C[1800,-1]:.4f} kg/kg。

题目给定参数下，热扩散率约为 $1.6886\\times10^{{-7}}\\ \\mathrm{{m^2/s}}$，初始水分扩散系数约为 $4.9377\\times10^{{-9}}\\ \\mathrm{{m^2/s}}$，前者约为后者的34.20倍。温度变化向内部传播较快，含水率变化更集中于表层，与这一尺度关系一致；随着表层含水率下降，局部扩散系数也减小。该解释对应本问热、质分别求解的模型，不涉及蒸发冷却对温度的反馈。

上述温差、含水率差和百分比由未舍入结果计算，因此与先对表中数值作四舍五入再求差可能相差末位。完整的逐秒、每0.1 cm结果保存于配套 `result1.xlsx`。'''
    mapping={'@CONVERGENCE_TABLE@':convergence,'@TABLE_T@':table(T),'@TABLE_C@':table(C),
             '@SERIES_TAIL@':f"{audit['heat_series_400_vs_800_max']:.3e}",
             '@HEAT_REFERENCE_ERROR@':f"{audit['runs']['bdf6400_128']['heat_series']['max_abs']:.3e}",
             '@HEAT_RESIDUAL@':f"{meta['heat_metrics'][0]:.3e}",
             '@WATER_RESIDUAL@':f"{meta['water_metrics'][0]:.3e}",
             '@PICARD@':str(meta['max_picard_iterations']),'@RESULT_ANALYSIS@':result_analysis}
    for k,v in mapping.items():numerical=numerical.replace(k,v)
    body=model+numerical
    assert '@TABLE' not in body and body.count('$$')%2==0
    (DEST/'问题一正文与分析.md').write_text(body)
    changes=[]
    for x in audit['supplied_excel_vs_supplied_paper']['T']['locations']:
        changes.append(f"| {x['time_s']} | {x['r_cm']:.1f} | {x['excel']:.4f} | {x['paper']:.4f} |")
    final=audit['runs']['bdf6400_128']
    report=f'''# 问题一核验结论与交接说明

## 使用的文件

写作手使用本目录的 `问题一正文与分析.md` 和 `result1.xlsx`。原始网页结果、原交接稿和题目模板均未覆盖。正文已补充“问题一的分析”、计算方法、两张结果表及结果解释。

## 核验发现

1. 已确定的三个物理假设、连续方程及初边值条件在约定范围内自洽，予以保留。水分通量需区分约化量与实际质量通量。
2. 从原始附件独立实现给定的N=1600、Δt=0.015625 s、后向欧拉、表面D(Cs)方案，原交接稿两张表的70个数在四位小数下全部复现。全量结果原始缓存为 `../独立核验/runs/be1600_64.npz`。
3. 用户后补的 `../result1.xlsx` 两表均为1800×21，时间和位置正确，无缺失；但所有结果单元格均为General格式，未固定四位小数显示。它与原论文温度表有以下5处不一致，证明Excel与文字结果未完全同步：

| 时间/s | r/cm | 上传Excel/℃ | 原交接稿/℃ |
| ---: | ---: | ---: | ---: |
{chr(10).join(changes)}

4. 原计算继续将时间步减半，论文温度35点中有4个四位小数变化；完整输出中温度有3673个、水分有79个变化。原先“最大差小于5×10⁻⁵即可保证四位小数不变”的说法不成立。
5. 单独构造圆柱导热特征函数级数解，验证温度场。原BE方案与级数解最大差约3.017×10⁻⁵ ℃；该误差足以改变靠近舍入阈值的末位。
6. 原守恒记录未提供相对残差的分母和近零处理，无法确认其数值定义。复算重新给出了绝对残差与明确归一化的残差；不能将原数值直接移植到新格式中。

## 本次交付采用的数值处理

三项物理假设保持不变。为减小末位误差，最终计算使用二阶后向差分BDF2（首步BE启动），N=6400、Δt=1/128 s、Picard容限10⁻¹⁰。内部界面仍采用调和平均；表面半单元采用圆柱几何阻力和扩散系数积分均值。它们是同一连续模型的数值改进，不能在正文中继续写成原来的“全程后向欧拉+仅取D(Cs)”。中心未知量明确按中心点值解释。

新正文已同步改写离散式、表面公式、守恒式和设置。相对于原交接稿，新论文温度表有 {final['vs_supplied_paper']['T']['round4_different']} 处末位修正，含水率表四位小数全部保持一致；相对于后补Excel，论文温度表有 {final['T']['round4_different_paper35']} 处不同。全Excel两场必须一起使用本次版本。

## 精度证据及其限度

{convergence}

空间细化最大含水率差发生在1 s表面，说明35个论文点不能覆盖最难的早期边界响应。上述空间和时间比较的70个论文数均保持四位小数一致。完整输出的空间比较仍有温度 {space['T']['round4_different_all']} 个、含水率 {space['C']['round4_different_all']} 个值跨舍入阈值；不能声称全75600个数已获得严格的四位真值保证。交付采用较细解，保留全部未舍入缓存供复核。

最终温度场与独立800项级数解最大差 {final['heat_series']['max_abs']:.6e} ℃；论文35点舍入一致。800项与400项差 {audit['heat_series_400_vs_800_max']:.6e} ℃。最终水分计算最多 {meta['max_picard_iterations']} 次Picard迭代；均衡初边值测试保持常量；全计算时间步未发现记录的单调性/径向方向性违反。

最终热量全局绝对残差最大值 {meta['heat_metrics'][0]:.6e} W，约化水分全局绝对残差最大值 {meta['water_metrics'][0]:.6e} m³/s（乘ρ_d成为kg水/s）。相对残差定义为 |Σstorage+A_R flux|/max(Σ|storage|,|A_R flux|,10⁻³⁰·2πL)，其中storage使用当前BE/BDF2对应的时间差分。最大相对全局残差分别为 {meta['heat_metrics'][1]:.6e} 和 {meta['water_metrics'][1]:.6e}。逐单元方程残差、表面残差也单独保存于运行JSON，不与全局残差混称。

这些检验支持当前模型下的数值一致性，不构成实测验证或对全部舍入末位的数学证明。

## 复现

在 `../独立核验` 目录运行 `/opt/anaconda3/bin/python reproduce.py` 可重新生成计算、比较和本目录交付物。需要numpy、scipy、numba、openpyxl。源文件 `verify_q1.py` 只读取原始附件，不修改原始数据。全部比较指标及差异位置见 `../独立核验/核验数据.json`。
'''
    (DEST/'核验结论与交接说明.md').write_text(report)
    files=[excel,DEST/'问题一正文与分析.md',DEST/'核验结论与交接说明.md',
           ROOT/'verify_q1.py',ROOT/'compare_q1.py',ROOT/'build_delivery.py',
           ROOT/'runs/bdf6400_128.npz',ROOT/'runs/bdf6400_128.json']
    manifest={'input_sha256':meta['input_sha256'],
              'final_settings':meta['settings'],'data_cells_per_sheet':37800,
              'numeric_output_format':'0.0000','paper_excel_consistency':'all 70 exact after rounding',
              'files':{str(p.relative_to(ROOT.parent)):sha(p) for p in files}}
    (DEST/'复现清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print('Excel shape, times, 75600 numeric values, all four-decimal formats and 70 paper cells: PASS')
    print('Delivery:',DEST)

if __name__=='__main__':main()
