"""Whole-workbook audit and independent Bessel-series heat reference."""
from pathlib import Path
import json, hashlib
import numpy as np
import openpyxl
from scipy.special import j0,j1,jn_zeros
from scipy.optimize import brentq
from verify_q1 import ROOT, INPUT, R, K, RHOCP, HT

PAPER_TIMES=np.array([100,300,600,900,1200,1500,1800])
PAPER_RADII=np.array([0,5,10,15,20])

def compare(a,b):
    diff=np.abs(a-b)
    ij=np.unravel_index(np.argmax(diff),diff.shape)
    rounded=np.round(a,4)!=np.round(b,4)
    pt=diff[PAPER_TIMES-1][:,PAPER_RADII]
    pr=rounded[PAPER_TIMES-1][:,PAPER_RADII]
    return {'max_abs':float(diff[ij]),'max_at_time_s':int(ij[0]+1),
            'max_at_r_cm':round(ij[1]/10,1),'round4_different_all':int(rounded.sum()),
            'max_abs_paper35':float(pt.max()),'round4_different_paper35':int(pr.sum()),
            'paper_changed':[{'t_s':int(PAPER_TIMES[i]),'r_cm':float(PAPER_RADII[j]/10),
               'a':float(a[PAPER_TIMES[i]-1,PAPER_RADII[j]]),
               'b':float(b[PAPER_TIMES[i]-1,PAPER_RADII[j]])}
               for i,j in zip(*np.where(pr))]}

def exact_heat(data,modes):
    bi=HT*R/K
    right=jn_zeros(0,modes)
    left=np.r_[0.,jn_zeros(1,modes-1)]
    roots=np.array([brentq(lambda x:x*j1(x)-bi*j0(x),a+1e-12,b)
                    for a,b in zip(left,right)])
    coef=2*j1(roots)/(roots*(j0(roots)**2+j1(roots)**2))
    lam=K/RHOCP*roots**2/R**2
    spatial=coef[None,:]*j0(np.arange(21)[:,None]/20*roots[None,:])
    modal=np.zeros(modes)
    output=np.empty((1801,21));output[0,:]=28
    decay=np.exp(-lam)
    for t in range(1,1801):
        idx=min((t-1)//60,29)
        slope=(data[idx+1,1]-data[idx,1])/60
        modal=modal*decay+slope*(-np.expm1(-lam))/lam
        ambient=data[idx,1]+slope*(t-data[idx,0])
        output[t,:]=ambient-spatial@modal
    return output

def main():
    src=ROOT.parent/'result1.xlsx'
    wb=openpyxl.load_workbook(src,data_only=True)
    supplied={};sheet_audit={}
    for name,key in [('温度','T'),('水分浓度','C')]:
        s=wb[name]; vals=list(s.values)
        a=np.array(vals[1:],dtype=float)
        assert a.shape==(1800,22) and np.array_equal(a[:,0],np.arange(1,1801))
        assert np.isfinite(a).all()
        assert np.allclose(np.array(vals[0][1:],float),np.arange(21)/10,atol=1e-12)
        supplied[key]=a[:,1:]
        sheet_audit[name]={'rows':s.max_row,'columns':s.max_column,
            'missing':0,'four_decimal_format_count':sum(c.number_format=='0.0000' for row in s.iter_rows(min_row=2,min_col=2) for c in row),
            'numeric_result_cells':37800}
    wb.close()
    data=np.array(list(openpyxl.load_workbook(INPUT,read_only=True,data_only=True).active.values)[1:32],float)
    t400=exact_heat(data,400);t800=exact_heat(data,800)
    np.savez_compressed(ROOT/'runs/heat_reference.npz',T=t800,input=data)
    audit={'supplied_file':str(src),'supplied_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),
           'sheets':sheet_audit,
           'heat_series_400_vs_800_max':float(np.max(np.abs(t400-t800))),
           'supplied_vs_heat_series':compare(supplied['T'],t800[1:]),'runs':{}}
    text=(ROOT.parent/'第一问计算结果与写作手交接_重算.md').read_text()
    table_rows=[]
    for line in text.splitlines():
        fields=[s.strip() for s in line.strip().strip('|').split('|')]
        if len(fields)==6 and fields[0] in {str(x) for x in PAPER_TIMES}:
            try:table_rows.append([float(x) for x in fields[1:]])
            except ValueError:pass
    assert len(table_rows)==14
    paper={'T':np.array(table_rows[:7]),'C':np.array(table_rows[7:])}
    audit['supplied_excel_vs_supplied_paper']={}
    for key in ['T','C']:
        selected=supplied[key][PAPER_TIMES-1][:,PAPER_RADII]
        mismatch=np.round(selected,4)!=np.round(paper[key],4)
        audit['supplied_excel_vs_supplied_paper'][key]={'different':int(mismatch.sum()),
          'locations':[{'time_s':int(PAPER_TIMES[i]),'r_cm':float(PAPER_RADII[j]/10),
             'excel':float(selected[i,j]),'paper':float(paper[key][i,j])} for i,j in zip(*np.where(mismatch))]}
    available={}
    for p in sorted((ROOT/'runs').glob('*.npz')):
        if p.stem=='heat_reference':continue
        z=np.load(p)
        if z['T'].shape!=(1801,21):continue
        available[p.stem]={'T':z['T'][1:],'C':z['C'][1:]}
        audit['runs'][p.stem]={key:compare(supplied[key],z[key][1:]) for key in ['T','C']}
        audit['runs'][p.stem]['heat_series']=compare(z['T'][1:],t800[1:])
        audit['runs'][p.stem]['vs_supplied_paper']={k:{
            'round4_different':int(np.sum(np.round(z[k][PAPER_TIMES][:,PAPER_RADII],4)!=paper[k])),
            'max_diff_from_printed':float(np.max(np.abs(z[k][PAPER_TIMES][:,PAPER_RADII]-paper[k])))} for k in ['T','C']}
    pairs=[('be1600_64','be1600_128'),('be1600_64','bdf1600_32'),
           ('bdf800_32','bdf1600_32'),('bdf1600_32','bdf1600_64'),
           ('bdf1600_64','bdf3200_64'),('be1600_64','be1600_integrated64'),
           ('bdf3200_64','bdf3200_128'),('bdf3200_128','bdf6400_128')]
    audit['pairs']={}
    for a,b in pairs:
        if a in available and b in available:
            audit['pairs'][a+'__'+b]={k:compare(available[a][k],available[b][k]) for k in ['T','C']}
    (ROOT/'核验数据.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
    compact={'xlsx':sheet_audit,'heat_reference_tail':audit['heat_series_400_vs_800_max'],
             'supplied_heat':audit['supplied_vs_heat_series'],
             'excel_paper':audit['supplied_excel_vs_supplied_paper'],
             'runs':{n:{k:{a:v for a,v in d.items() if a!='paper_changed'} for k,d in r.items()} for n,r in audit['runs'].items()},
             'pairs':{n:{k:{a:v for a,v in d.items() if a!='paper_changed'} for k,d in r.items()} for n,r in audit['pairs'].items()}}
    print(json.dumps(compact,ensure_ascii=False,indent=2))

if __name__=='__main__':main()
