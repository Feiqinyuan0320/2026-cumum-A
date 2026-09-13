"""Create Q4 deliverables from computed and hashed case archives."""
from pathlib import Path
import argparse
import csv
import json
import sys
import numpy as np
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import solve_q4 as q

ROOT = q.ROOT
RESULTS = ROOT/'results'


def load(label):
    meta = json.loads((ROOT/'runs'/f'{label}.json').read_text())
    assert meta['solver_sha256'] == q.digest(q.__file__)
    assert meta['environment_sha256'] == q.digest(q.ENV)
    assert meta['radius_input_sha256'] == q.digest(q.RAD)
    with np.load(ROOT/'runs'/f'{label}.npz') as f:
        a = dict(f)
    assert np.all(np.diff(a['times']) == 60) and a['times'][0] == 60
    assert np.allclose(a['event'], meta['event'], rtol=0, atol=0)
    assert a['event'][3] < .15 and a['times'][-1] >= a['event'][2]
    return a, meta


def workbook(a):
    wb = openpyxl.load_workbook(ROOT.parent/'附件/结果文件夹/result4.xlsx')
    sheet = wb.active
    for row in sheet:
        for cell in row:
            cell.value = None
    header = ['时间/s\\到药材中心的距离/cm', *[i/10 for i in range(20)], '药材表面']
    for j, name in enumerate(header, 1):
        cell = sheet.cell(1, j, name)
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='23475B')
        cell.alignment = Alignment(horizontal='center')
    for i, (t, values) in enumerate(zip(a['times'], a['C']), 2):
        sheet.cell(i, 1, int(t))
        for j, value in enumerate(values, 2):
            cell = sheet.cell(i, j)
            cell.value = float(np.round(value, 4)) if np.isfinite(value) else None
            cell.number_format = '0.0000'
    sheet.freeze_panes = 'B2'
    sheet.column_dimensions['A'].width = 34
    for j in range(2, 23):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 11
    path = RESULTS/'result4.xlsx'
    wb.save(path); wb.close()
    # Verify the exported workbook itself, including domain blanks and surface equality.
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.values)
    assert list(rows[0]) == header and len(rows) == len(a['times'])+1
    values = np.asarray(rows[1:], dtype=float)
    assert values.shape == (len(a['times']), 22)
    assert np.array_equal(values[:, 0], a['times'])
    assert np.array_equal(values[:, 1:], np.round(a['C'], 4), equal_nan=True)
    assert all(c.number_format == '0.0000' for row in wb.active.iter_rows(min_row=2, min_col=2) for c in row)
    wb.close()
    return int(np.isnan(a['C']).sum())


def table6(a):
    times = np.arange(21600., a['event'][2], 21600.)
    idx = np.searchsorted(a['times'], times)
    assert np.array_equal(a['times'][idx], times)
    positions = [0, 5, 10, 20]
    water = np.vstack((a['C'][idx][:, positions], a['event_C'][positions]))
    hours = np.r_[times/3600, a['event'][2]/3600]
    radii = np.r_[a['radii'][idx], a['event_radius']]*100
    assert np.isfinite(water).all() and np.all(radii < 1.5)
    header = ['时间/h', '0 cm', '0.5 cm', '1 cm', '药材表面']
    wb = openpyxl.Workbook(); sh = wb.active; sh.title = '表6'
    sh.append(header)
    for h, values in zip(hours, water):
        sh.append([float(h), *np.round(values, 4).tolist()])
    for row in sh.iter_rows(min_row=2):
        for c in row:
            c.number_format = '0.0000'
    wb.save(RESULTS/'表6_水分浓度.xlsx'); wb.close()
    lines = ['| '+' | '.join(header)+' |', '| ---: | ---: | ---: | ---: | ---: |']
    texrows = []
    for i, (h, values) in enumerate(zip(hours, water)):
        label = f'{h:.4f}（结束）' if i == len(hours)-1 else f'{h:g}'
        lines.append('| '+label+' | '+' | '.join(f'{v:.4f}' for v in values)+' |')
        texrows.append(label+' & '+' & '.join(f'{v:.4f}' for v in values)+r' \\')
    (RESULTS/'表6_水分浓度.md').write_text('\n'.join(lines)+'\n')
    tex = [r'\begin{table}[htbp]', r'\centering', r'\caption{考虑尺寸收缩时药材各位置的干基含水率（kg/kg）}',
           r'\label{tab:q4-moisture}', r'\begin{tabular}{rrrrr}', r'\toprule',
           r'时间/h & 0 cm & 0.5 cm & 1 cm & 药材表面 \\', r'\midrule',
           *texrows, r'\bottomrule', r'\end{tabular}', r'\end{table}']
    (RESULTS/'表6_水分浓度.tex').write_text('\n'.join(tex)+'\n')
    return hours, water, radii, lines


def cases():
    out = []
    for path in sorted((ROOT/'runs').glob('*.json')):
        m = json.loads(path.read_text())
        if m.get('event') is None or m.get('solver_sha256') != q.digest(q.__file__):
            continue
        s, e = m['settings'], m['event']
        out.append({'label': path.stem, 'geometry': s['geometry'], 'N': s['n'],
                    'early_dt_s': s['early_dt'], 'dt_s': s['dt'],
                    'crossing_s': (e[0]+e[1])/2, 'crossing_h': (e[0]+e[1])/7200,
                    'reported_h': e[2]/3600, 'event_interval_s': e[1]-e[0]})
    return out


def plot(a, b, summary):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from utils.plot_style import apply_publication_style, export_figure, COLOR_SEQUENCE
    apply_publication_style(language='en')
    plt.rcParams.update({'font.sans-serif': ['DejaVu Sans'], 'font.size': 9,
                         'axes.labelsize': 10, 'xtick.labelsize': 9, 'ytick.labelsize': 9,
                         'legend.fontsize': 9})
    figdir = ROOT/'figures'
    figdir.mkdir(exist_ok=True)
    def save(fig, name):
        export_figure(fig, figdir/name, dpi=300, grayscale_preview=True)
        fig.savefig(figdir/f'{name}.pdf')
        plt.close(fig)
    r = a['radius_data']
    fig, ax = plt.subplots(figsize=(6.3, 3.1), layout='constrained')
    ax.plot(r[:, 0]/3600, r[:, 1]*100, color=COLOR_SEQUENCE[0], lw=1.2, label='Piecewise linear radius')
    ax.scatter(r[:, 0]/3600, r[:, 1]*100, s=9, facecolors='white', edgecolors=COLOR_SEQUENCE[0],
               linewidths=.55, label='Measured radius', zorder=3)
    ax.set(xlabel='Drying time (h)', ylabel='Specimen radius (cm)', xlim=(0, 72), ylim=(1.15, 2.05))
    ax.legend(loc='upper right')
    save(fig, 'raw_q4_实测半径与插值')
    fig, ax = plt.subplots(figsize=(6.3, 3.6), layout='constrained')
    styles = ['-', '--', '-.', ':', (0, (5, 2, 1, 2))]
    for j in range(5):
        ax.plot(np.r_[0, a['times']/3600], np.r_[2.55, a['relative_C'][:, j]],
                color=COLOR_SEQUENCE[j], ls=styles[j], lw=1.5, label=f'r/R(t) = {j/4:g}')
    ax.axhline(.15, color='#444444', lw=.8, ls='--')
    ax.set(xlabel='Drying time (h)', ylabel='Dry-basis moisture content (kg/kg)',
           xlim=(0, a['times'][-1]/3600), ylim=(0, 2.65))
    ax.legend(loc='upper right')
    save(fig, 'process_q4_材料位置含水率')
    fig, ax = plt.subplots(figsize=(6.3, 2.6), layout='constrained')
    durations = [b['event'][2]/3600, a['event'][2]/3600]
    bars = ax.barh([1, 0], durations, height=.48, color=[COLOR_SEQUENCE[1], COLOR_SEQUENCE[0]])
    ax.set_yticks([1, 0], ['Fixed radius (2 cm)', 'Measured shrinkage'])
    ax.set(xlabel='Drying time (h)', xlim=(0, max(durations)*1.17), ylim=(-.6, 1.7))
    for bar, t in zip(bars, durations):
        ax.text(t+max(durations)*.02, bar.get_y()+bar.get_height()/2, f'{t:.2f} h', va='center', fontsize=9)
    ax.text(.98, .94, f'Same material properties; time reduction {summary["reduction_percent"]:.2f}%',
            transform=ax.transAxes, ha='right', va='top', fontsize=9)
    ax.grid(axis='y', visible=False)
    save(fig, 'result_q4_收缩与固定半径比较')


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--main', required=True)
    p.add_argument('--fixed', required=True)
    p.add_argument('--no-plots', action='store_true')
    args = p.parse_args()
    a, ma = load(args.main); b, mb = load(args.fixed)
    assert ma['geometry'] == 'shrink' and mb['geometry'] == 'fixed'
    for key in ['n', 'dt', 'early_dt', 'extension']:
        assert ma['settings'][key] == mb['settings'][key], key
    RESULTS.mkdir(exist_ok=True)
    blanks = workbook(a)
    hours, water, radii, lines = table6(a)
    cc = cases()
    with (RESULTS/'计算方案比较.csv').open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(cc[0])); w.writeheader(); w.writerows(cc)
    s = {'main_case': args.main, 'fixed_case': args.fixed, 'main_settings': ma['settings'],
         'event_main': ma['event'], 'event_fixed': mb['event'],
         'reported_drying_h': a['event'][2]/3600, 'fixed_drying_h': b['event'][2]/3600,
         'reduction_h': (b['event'][2]-a['event'][2])/3600,
         'reduction_percent': (1-a['event'][2]/b['event'][2])*100,
         'event_radius_cm': float(a['event_radius']*100), 'boundary_mean_T_C': ma['tail_T_C'],
         'radius_plateau_extension_used': ma['radius_plateau_extension_used'],
         'result4_rows': len(a['times']), 'result4_columns': 22,
         'outside_domain_blank_cells': blanks, 'result4_last_s': float(a['times'][-1]),
         'table6_hours': hours.tolist(), 'table6_unrounded_C': water.tolist(),
         'table6_radius_cm': radii.tolist(), 'all_output_maximum_at_axis': bool(np.all(a['maxima'][:, 1] == 0)),
         'main_stats': ma['stats'], 'fixed_stats': mb['stats'], 'cases': cc}
    (RESULTS/'计算摘要.json').write_text(json.dumps(s, ensure_ascii=False, indent=2))
    with (RESULTS/'绘图数据.csv').open('w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['time_h', 'radius_cm', 'C_xi0', 'C_xi025', 'C_xi05', 'C_xi075', 'C_surface', 'C_max'])
        w.writerows(np.column_stack((a['times']/3600, a['radii']*100, a['relative_C'], a['maxima'][:, 0])))
    sys.path.insert(0, '/Users/baitutu/.codex/skills/math-modeling/tools/figure/scripts')
    from profile_data import profile_data
    profile = profile_data(str(RESULTS/'绘图数据.csv'))
    (RESULTS/'绘图数据剖析.json').write_text(json.dumps(profile, ensure_ascii=False, indent=2, default=str))
    contract = {'scope': 'q4; 按用户节省额度及已确认方案，只生成3张核心图。',
        'backend': 'matplotlib', 'formats': ['PDF', 'SVG editable text', 'PNG 300 dpi', 'grayscale PNG'],
        'statistics': '确定性模型与题给半径数据，无抽样置信区间。材料位置曲线不等于固定物理半径曲线。',
        'profile': {'radius_rows': len(a['radius_data']), 'missing_radius': 0,
                    'relative_output_rows': len(a['times']), 'missing_relative_C': 0,
                    'outside_domain_blanks': blanks, 'units': 'h, cm, kg/kg dry basis'},
        'figures': [
            {'name': 'raw_q4_实测半径与插值', 'claim': '实测半径在前期快速下降，随后接近1.2 cm。',
             'evidence': '145个真实半径节点与逐段直线', 'layout': 'single axis scatter and line, 6.3 x 3.1 in'},
            {'name': 'process_q4_材料位置含水率', 'claim': '固定材料位置的失水进程呈径向差异。',
             'evidence': 'xi=0,.25,.5,.75,1的求解输出', 'layout': 'single axis lines, 6.3 x 3.6 in'},
            {'name': 'result_q4_收缩与固定半径比较', 'claim': f'相同物性下收缩使预测时长缩短{s["reduction_percent"]:.2f}%。',
             'evidence': '相同数值设置的两组全程独立计算', 'layout': 'zero-baseline horizontal bars, 6.3 x 2.6 in'}]}
    (RESULTS/'图表契约.json').write_text(json.dumps(contract, ensure_ascii=False, indent=2))
    print(json.dumps({k: s[k] for k in ['reported_drying_h', 'fixed_drying_h', 'reduction_percent',
                                       'result4_rows', 'outside_domain_blank_cells']}, ensure_ascii=False))
    print('Workbook read-back PASS')
    if not args.no_plots:
        plot(a, b, s)


if __name__ == '__main__':
    main()
