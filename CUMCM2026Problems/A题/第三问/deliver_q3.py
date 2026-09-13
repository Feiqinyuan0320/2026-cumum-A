"""Export Q3 tables, evidence summary and three core figures from saved runs only."""
from pathlib import Path
from copy import copy
import argparse
import csv
import json
import sys

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
SKILL = Path('/Users/baitutu/.codex/skills/math-modeling')


def load(label):
    with np.load(ROOT/'runs'/f'{label}.npz') as a:
        out = dict(a)
    meta = json.loads((ROOT/'runs'/f'{label}.json').read_text())
    return out, meta


def spreadsheet(path, times, values):
    template = ROOT.parent/'附件/结果文件夹/result3.xlsx'
    wb = openpyxl.load_workbook(template)
    sheet = wb.active
    # Original template has only a header and three example times plus ellipses.
    for row in sheet:
        for cell in row:
            cell.value = None
    sheet.cell(1, 1, '时间/s\\到药材中心的距离/cm')
    for j in range(21):
        sheet.cell(1, j+2, j/10)
    for i, (t, row) in enumerate(zip(times, values), 2):
        sheet.cell(i, 1, int(round(t)))
        for j, x in enumerate(row, 2):
            cell = sheet.cell(i, j, float(np.round(x, 4)))
            cell.number_format = '0.0000'
    sheet.freeze_panes = 'B2'
    sheet.column_dimensions['A'].width = 32
    for cell in sheet[1]:
        cell.font = Font(name='Arial', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='23475B')
        cell.alignment = Alignment(horizontal='center')
    for j in range(2, 23):
        sheet.column_dimensions[openpyxl.utils.get_column_letter(j)].width = 11
    wb.save(path)
    wb.close()


def main():
    p = argparse.ArgumentParser()
    p.add_argument('--main', default='mean_n3200_dt0p5')
    p.add_argument('--last', default='last_n3200_dt0p5')
    p.add_argument('--no-plots', action='store_true')
    args = p.parse_args()
    a, meta = load(args.main)
    b, last = load(args.last)
    results = ROOT/'results'
    results.mkdir(exist_ok=True)
    assert np.all(np.diff(a['times']) == 60) and a['times'][0] == 60
    assert a['C'].shape == (len(a['times']), 21) and np.isfinite(a['C']).all()
    assert a['event'][3] < 0.15 and a['times'][-1] >= a['event'][2]
    spreadsheet(results/'result3.xlsx', a['times'], a['C'])

    times = np.arange(21600.0, a['event'][2], 21600.0)
    indices = np.searchsorted(a['times'], times)
    assert np.all(a['times'][indices] == times)
    water = np.vstack((a['C'][indices][:, ::5], a['event_C'][::5]))
    hours = np.r_[times/3600, a['event'][2]/3600]
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = '表5'
    s.append(['时间/h', '0 cm', '0.5 cm', '1 cm', '1.5 cm', '2 cm', '说明'])
    for i, (h, row) in enumerate(zip(hours, water)):
        s.append([float(h), *np.round(row, 4).tolist(), '烘干结束' if i == len(hours)-1 else ''])
    for row in s.iter_rows(min_row=2, max_col=6):
        for c in row:
            c.number_format = '0.0000'
    s.freeze_panes = 'B2'
    wb.save(results/'表5_水分浓度.xlsx')
    wb.close()

    lines = ['| 时间/h | 0 cm | 0.5 cm | 1 cm | 1.5 cm | 2 cm |',
             '| ---: | ---: | ---: | ---: | ---: | ---: |']
    for i, (h, row) in enumerate(zip(hours, water)):
        label = f'{h:.4f}（结束）' if i == len(hours)-1 else f'{h:g}'
        lines.append('| '+label+' | '+' | '.join(f'{x:.4f}' for x in row)+' |')
    (results/'表5_水分浓度.md').write_text('\n'.join(lines)+'\n')
    tex_rows = []
    for i, (h, row) in enumerate(zip(hours, water)):
        label = f'{h:.4f}（结束）' if i == len(hours)-1 else f'{h:g}'
        tex_rows.append(label+' & '+' & '.join(f'{x:.4f}' for x in row)+r' \\')
    (results/'表5_水分浓度.tex').write_text('\n'.join([
        r'\begin{table}[htbp]', r'\centering', r'\caption{药材烘干过程的水分浓度（kg/kg，干基）}',
        r'\label{tab:q3-moisture}', r'\begin{tabular}{rrrrrr}', r'\toprule',
        r'时间/h & 0 cm & 0.5 cm & 1 cm & 1.5 cm & 2 cm \\', r'\midrule',
        *tex_rows, r'\bottomrule', r'\end{tabular}', r'\end{table}', '']))

    cases = []
    for file in sorted((ROOT/'runs').glob('*.json')):
        if not file.stem.startswith(('mean_n', 'last_n')):
            continue
        d = json.loads(file.read_text())
        if d.get('event') and d.get('settings', {}).get('mode') == 'case':
            e = d['event']
            cases.append({'label': file.stem, 'N': d['settings']['n'], 'dt_s': d['settings']['dt'],
                          'extension': d['settings']['extension'], 'crossing_s': (e[0]+e[1])/2,
                          'crossing_h': (e[0]+e[1])/7200, 'reported_h': e[2]/3600,
                          'event_bracket_width_s': e[1]-e[0], 'max_C_at_report': e[3],
                          'max_iterations': d['stats'][0]})
    with (results/'计算方案比较.csv').open('w', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(cases[0]))
        w.writeheader(); w.writerows(cases)
    early = json.loads((ROOT/'runs/q2_regression.json').read_text())
    summary = {
        'main_case': args.main, 'comparison_case': args.last,
        'main_settings': meta['settings'], 'boundary_mean_T_C': meta['tail_T_C'],
        'boundary_last_T_C': last['tail_T_C'],
        'event_main': meta['event'], 'event_last': last['event'],
        'reported_drying_h': a['event'][2]/3600,
        'boundary_effect_seconds_last_minus_mean': np.mean(b['event'][:2])-np.mean(a['event'][:2]),
        'boundary_effect_percent': (np.mean(b['event'][:2])/np.mean(a['event'][:2])-1)*100,
        'result3_rows': len(a['times']), 'result3_columns': 22,
        'result3_last_s': float(a['times'][-1]),
        'table5_hours': hours.tolist(), 'table5_unrounded_C': water.tolist(),
        'all_output_maximum_at_axis': bool(np.all(a['maxima'][:, 1] == 0)),
        'long_time_solver_stats': meta['stats'], 'q2_regression': early['regression'], 'cases': cases,
    }
    (results/'计算摘要.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    write_analysis(summary, lines)
    with (results/'绘图数据.csv').open('w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['time_h', 'C_r0', 'C_r0p5', 'C_r1', 'C_r1p5', 'C_r2', 'C_max'])
        w.writerows(np.column_stack((a['times']/3600, a['C'][:, ::5], a['maxima'][:, 0])))

    # Read back the actual workbook: exact output grid, numeric cells and rounding.
    wb = openpyxl.load_workbook(results/'result3.xlsx', data_only=True, read_only=True)
    rows = list(wb.active.values)
    actual = np.asarray(rows[1:], dtype=float)
    assert actual.shape == (len(a['times']), 22)
    assert np.array_equal(actual[:, 0], a['times'])
    assert np.array_equal(actual[:, 1:], np.round(a['C'], 4))
    assert all(cell.number_format == '0.0000' for row in wb.active.iter_rows(min_row=2, min_col=2) for cell in row)
    wb.close()
    print(json.dumps({'reported_drying_h': summary['reported_drying_h'],
                      'result3_rows': len(a['times']), 'boundary_effect_s': summary['boundary_effect_seconds_last_minus_mean'],
                      'workbook_readback': 'PASS'}, ensure_ascii=False))
    if not args.no_plots:
        plot(a, b, cases, args.main, args.last)


def write_analysis(s, paper_table):
    e, last = s['event_main'], s['event_last']
    hours = s['reported_drying_h']
    cases = sorted(s['cases'], key=lambda c: (c['extension'] != 'mean', c['N'], -c['dt_s']))
    compare = ['| 边界延拓 | 网格数 N | 3 h后步长/s | 阈值交点时间/h |',
               '| --- | ---: | ---: | ---: |']
    for c in cases:
        compare.append(f'| {"尾段均值" if c["extension"] == "mean" else "末值"} | {c["N"]} | {c["dt_s"]:g} | {c["crossing_h"]:.6f} |')
    mean_cases = [c for c in cases if c['extension'] == 'mean' and c['dt_s'] == s['main_settings']['dt']]
    mean_cases.sort(key=lambda c: c['N'])
    spatial = ''
    if len(mean_cases) > 1:
        low, high = mean_cases[-2:]
        spatial = f'在相同步长下，最近两级网格 N={low["N"]} 与 N={high["N"]} 的结束时间相差 {abs(low["crossing_s"]-high["crossing_s"]):.4f} s。'
    time_n = s['main_settings']['n']
    time_cases = sorted([c for c in cases if c['extension'] == 'mean' and c['N'] == time_n], key=lambda c: -c['dt_s'])
    temporal = ''
    if len(time_cases) > 1:
        aa, bb = time_cases[-2:]
        temporal = f'在 N={time_n} 下，步长由 {aa["dt_s"]:g} s 减为 {bb["dt_s"]:g} s 时，结束时间相差 {abs(aa["crossing_s"]-bb["crossing_s"]):.4f} s。'
    markdown = fr'''# 第三问：烘干时长的确定与结果分析

## 问题分析与长期边界

第三问沿用第二问的固定几何、变物性径向热质传递模型，需要将计算延续到药材各处干基含水率均低于0.15 kg/kg。新增的问题是4 h以后的环境输入以及首次全域达标时间的确定。采用平均含水率作为终止指标会掩盖内部尚未干燥的区域，因此以整个径向计算域上的最大含水率作为判断依据。

附件1提供0—4 h的环境数据，前4 h按采样点作分段线性插值。3—4 h的环境已接近稳定，本问以该时段61个节点的算术均值作为后续恒温干燥阶段的输入：烘房温度取{s['boundary_mean_T_C'][0]:.8f} ℃，等效外部水分浓度取{s['boundary_mean_T_C'][1]:.10f} kg/kg。该延拓属于模型假设，不是附件提供的未来观测；保留末值延拓作为对照。

## 全域达标判据与求解

定义 $C_{{\max}}(t)=\max_{{0\le r\le R}}C(r,t)$，以 $C_{{\max}}(t)<0.15$ 为停止条件。程序每步检查全部单元、轴线重构值与表面值，未预设中心必然最后达标。

数值方法保持第二问的圆柱有限体积、后向欧拉与耦合Picard迭代。扩散系数使用药材局部绝对温度，表面状态由Robin边界和半单元阻力反算。前3 h使用0.0625 s步长，最终采用 N={s['main_settings']['n']}、3 h以后步长{s['main_settings']['dt']:g} s。更细网格从原始均匀初态计算，未由稀疏的径向结果表替代内部状态。

发现首次跨越阈值的时间步后，从该步的前一状态起用后向欧拉子步二分定位，将事件区间收紧至0.001 s以内。连续过程严格不等式集合的下确界对应阈值交点，交点本身通常满足等号；为给出数值上严格达标的有限时刻，将上界对应的小时数向上取至四位小数，并重新计算该报告时刻。

## 计算结果

在上述边界和模型假设下，烘干所需时间约为 **{hours:.4f} h**，即约 **{hours/24:.4f} d**。未取整的阈值区间为 **[{e[0]:.9f}, {e[1]:.9f}] s**。在报告时刻，全域最大干基含水率为 **{e[3]:.12f} kg/kg**，小于0.15；最大值位于距轴线 **{e[4]*100:g} cm** 处。此处关于最后达标位置的判断来自实际计算。

表5给出每6 h及结束时刻的干基含水率：

{chr(10).join(paper_table)}

温度场在前期升温后趋于稳定，水分场仍保持从内部到表面的径向差异。中心含水率下降较慢，表层更早失水。随含水率降低，附录3中的 $\exp(-0.45/C)$ 减小，局部有效扩散能力下降，后期失水过程因而放缓。该解释对应本问有效扩散模型，不据此排除表面对流阻力的影响。

表中结束时的最大值按四位小数可能显示为0.1500；达标判断使用未舍入值，不能据显示值反推未达标。完整结果见 `results/result3.xlsx`，含{s['result3_rows']}个时间点、21个半径位置，从60 s起每60 s输出。最后一个整分钟为{s['result3_last_s']:.0f} s，覆盖精确终止时刻；精确结束行另列于表5，不插入逐分钟序列。

## 数值收敛与边界延拓影响

恢复的求解器在 N=400、步长0.0625 s下复现第二问0—3 h的温度和含水率，全部453600个输出数在四位小数下与已有result2.xlsx一致。该比较说明接续口径一致，不是相对于实测真值的误差检验。

各组计算的阈值交点时间如下，未用向上取整后的报告时间掩盖细小的离散差异：

{chr(10).join(compare)}

{temporal}{spatial}事件定位的毫秒级区间仅描述最后一步内的定位分辨率，不代表偏微分方程解具有毫秒级总体精度。四位小数按题面用于输出格式，当前网格比较仍足以影响若干末位，不宣称所有输出都已得到四位真值保证。

保持相同网格和步长，将4 h以后改为末值延拓，环境为{s['boundary_last_T_C'][0]:g} ℃、{s['boundary_last_T_C'][1]:g} kg/kg，报告时长为{last[2]/3600:.4f} h。其阈值交点时间比均值方案变化{s['boundary_effect_seconds_last_minus_mean']:.4f} s，相对变化{s['boundary_effect_percent']:.4f}%。该差异反映两项明确边界假设对预测的影响，不能作为所有未来环境波动造成误差的上界。

## 适用范围

本结果沿用固定尺寸、中段一维径向、宏观有效物性和外界等效浓度等假设，未引入尺寸收缩或未闭合的潜热项。数值收敛与第二问回归支持计算的一致性；由于没有内部温度、含水率和实际烘干终点的测量数据，不能将其称为实验验证。长期边界及物性经验式仍是预测依赖的关键条件。

## 图的使用

- `figures/result_q3_代表半径含水率.pdf`：展示各代表半径的完整失水过程。
- `figures/process_q3_全域达标与边界延拓.pdf`：展示阈值附近的最大含水率及两种延拓的交点。
- `figures/process_q3_结束时间收敛比较.pdf`：展示不同网格与步长的结束时间差，不含统计置信区间。

本问沿用第二问已核验的方法参考：Da Silva等（2014），DOI 10.1016/j.jfoodeng.2014.05.010；中文背景参考王乐意等（2024），DOI 10.11975/j.issn.1002-6819.202306104。参考文献随整篇统一编号，新增的边界均值延拓由本题附件与所述假设给出，不归于上述论文的实测结论。
'''
    (ROOT/'第三问结果与分析.md').write_text(markdown)


def plot(a, b, cases, main_label, last_label):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from utils.plot_style import apply_publication_style, export_figure, COLOR_SEQUENCE
    sys.path.insert(0, str(SKILL/'tools/figure/scripts'))
    from profile_data import profile_data
    profile = profile_data(str(ROOT/'results/绘图数据.csv'))
    (ROOT/'results/绘图数据剖析.json').write_text(json.dumps(profile, ensure_ascii=False, indent=2, default=str))
    apply_publication_style(language='en')
    plt.rcParams.update({'font.sans-serif': ['DejaVu Sans'], 'font.size': 9, 'axes.labelsize': 10, 'xtick.labelsize': 9,
                         'ytick.labelsize': 9, 'legend.fontsize': 9})
    figures = ROOT/'figures'

    def save(fig, name):
        export_figure(fig, figures/name, dpi=300, grayscale_preview=True)
        fig.savefig(figures/f'{name}.pdf')
        plt.close(fig)

    fig, ax = plt.subplots(figsize=(6.3, 3.7), layout='constrained')
    x = np.r_[0, a['times']/3600]
    styles = ['-', '--', '-.', ':', (0, (5, 2, 1, 2))]
    for j, radius in enumerate([0, 0.5, 1, 1.5, 2]):
        ax.plot(x, np.r_[2.55, a['C'][:, j*5]], label=f'r = {radius:g} cm',
                color=COLOR_SEQUENCE[j], linestyle=styles[j], lw=1.45)
    ax.axhline(0.15, color='#303030', ls='--', lw=0.9)
    ax.set(xlabel='Drying time (h)', ylabel='Dry-basis moisture content (kg/kg)', xlim=(0, x[-1]), ylim=(0, 2.65))
    ax.legend(loc='upper right')
    save(fig, 'result_q3_代表半径含水率')

    fig, ax = plt.subplots(figsize=(6.3, 3.3), layout='constrained')
    for out, label, color, style in [(a, 'Last-hour mean extension', COLOR_SEQUENCE[0], '-'),
                                      (b, 'Last-value extension', COLOR_SEQUENCE[1], '--')]:
        ax.plot(out['times']/3600, out['maxima'][:, 0], label=label, color=color, ls=style, lw=1.5)
        ax.plot(out['event'][2]/3600, out['event'][3], 'o', color=color, ms=4)
    left = min(a['event'][2], b['event'][2])/3600-0.6
    right = max(a['event'][2], b['event'][2])/3600+0.1
    ax.axhline(0.15, color='#303030', ls=':', label='Threshold: 0.15')
    ax.set(xlim=(left, right), ylim=(0.1497, 0.151), xlabel='Drying time (h)',
           ylabel='Maximum moisture content\n(dry basis, kg/kg)')
    ax.ticklabel_format(axis='y', style='plain', useOffset=False)
    ax.legend(loc='upper right')
    save(fig, 'process_q3_全域达标与边界延拓')

    numeric = sorted([c for c in cases if c['extension'] == 'mean'], key=lambda c: (c['N'], -c['dt_s']))
    base = next(c['crossing_s'] for c in cases if c['label'] == main_label)
    fig, ax = plt.subplots(figsize=(6.3, 3.5), layout='constrained')
    yy = np.arange(len(numeric))
    delta = [c['crossing_s']-base for c in numeric]
    ax.scatter(delta, yy, color=COLOR_SEQUENCE[0], s=32, zorder=3)
    ax.axvline(0, color='#555555', ls=':', lw=0.9)
    ax.set_yticks(yy, [f'N={c["N"]}, Δt={c["dt_s"]:g} s' for c in numeric])
    ax.invert_yaxis()
    ax.set(xlabel='Threshold time shift from adopted setting (s)', ylabel='Numerical settings')
    save(fig, 'process_q3_结束时间收敛比较')


if __name__ == '__main__':
    main()
